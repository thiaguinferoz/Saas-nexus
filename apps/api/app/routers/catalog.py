import csv
import io
import re
import unicodedata
import uuid
from decimal import Decimal, InvalidOperation
from numbers import Number

from fastapi import APIRouter, File, Form, HTTPException, Response, UploadFile, status
from sqlalchemy import delete, select, update
from sqlalchemy.exc import IntegrityError

from app.dependencies import CsrfGuard, CurrentTenant, DbSession
from app.models import CatalogItem
from app.schemas import CatalogImportRead, CatalogItemCreate, CatalogItemRead, CatalogItemUpdate

router = APIRouter(prefix="/catalog", tags=["catalog"])

MAX_CATALOG_FILE_BYTES = 20 * 1024 * 1024
MAX_CATALOG_ROWS = 1_000

CATALOG_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": (
        "nome", "produto", "servico", "item", "descricao", "nomedoproduto",
        "nomedoservico", "nomeproduto", "nomeservico", "descricaoitem",
        "produtoouservico",
    ),
    "price": (
        "preco", "valor", "price", "valorunitario", "precounitario", "precovenda",
        "precodevenda", "valorvenda", "valordevenda", "valorproduto", "valorservico",
        "precoproduto", "precoservico", "valorr", "valorrs", "valorbrl", "vlr",
        "vlrunitario", "custounitario", "custo",
    ),
    "category": ("categoria", "tipo", "category", "grupo", "departamento"),
    "sku": (
        "sku", "codigo", "cod", "idproduto", "codigoproduto", "codigodoproduto",
        "referencia", "ref",
    ),
    "description": (
        "detalhes", "observacao", "observacoes", "descricaocompleta", "detalhamento",
    ),
    "currency": ("moeda", "currency"),
}


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def first_value(row: dict[str, object], field: str) -> object | None:
    for alias in CATALOG_HEADER_ALIASES[field]:
        value = row.get(alias)
        if value is not None and str(value).strip():
            return value
    for header, value in row.items():
        if value is None or not str(value).strip():
            continue
        if any(len(alias) >= 4 and alias in header for alias in CATALOG_HEADER_ALIASES[field]):
            return value
    return None


def matching_header(headers: list[str], field: str) -> str | None:
    aliases = CATALOG_HEADER_ALIASES[field]
    for header in headers:
        if normalize_header(header) in aliases:
            return header
    for header in headers:
        normalized = normalize_header(header)
        if any(len(alias) >= 4 and alias in normalized for alias in aliases):
            return header
    return None


def parse_price(value: object | None) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Number) and not isinstance(value, bool):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None

    number = re.sub(r"[^0-9,.-]", "", str(value).strip())
    if not number:
        return None
    if "," in number and "." in number:
        decimal_separator = "," if number.rfind(",") > number.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        number = number.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in number or "." in number:
        separator = "," if "," in number else "."
        occurrences = number.count(separator)
        decimal_places = len(number.rsplit(separator, 1)[-1])
        if decimal_places in (1, 2):
            number = number.replace(separator, "", occurrences - 1).replace(separator, ".")
        else:
            number = number.replace(separator, "")
    try:
        return Decimal(number).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def header_score(values: tuple[object, ...] | list[object]) -> int:
    headers = [str(value) for value in values if value is not None]
    return sum(matching_header(headers, field) is not None for field in CATALOG_HEADER_ALIASES)


def tabular_rows(values: list[tuple[object, ...]]) -> tuple[int, list[tuple[int, dict[str, object]]]]:
    if not values:
        return 0, []
    candidates = [(header_score(row), index) for index, row in enumerate(values[:20])]
    score, header_index = max(candidates, default=(0, 0))
    if score == 0:
        populated = [
            (sum(value is not None and str(value).strip() != "" for value in row), index)
            for index, row in enumerate(values[:20])
        ]
        column_count, header_index = max(populated, default=(0, 0))
        if column_count < 2:
            return 0, []
    headers = [str(header).strip() if header is not None and str(header).strip() else f"Coluna {index + 1}" for index, header in enumerate(values[header_index])]
    rows = [
        (
            row_index + 1,
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)},
        )
        for row_index, row in enumerate(values[header_index + 1 :], start=header_index + 1)
        if any(value is not None and str(value).strip() for value in row)
    ]
    return score, rows


def spreadsheet_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, object]]]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        matrix = [tuple(row) for row in csv.reader(io.StringIO(text), dialect)]
        _, rows = tabular_rows(matrix)
        return rows
    if suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        best_score = 0
        best_rows: list[tuple[int, dict[str, object]]] = []
        for sheet in workbook.worksheets:
            matrix = list(sheet.iter_rows(max_row=MAX_CATALOG_ROWS + 20, values_only=True))
            score, rows = tabular_rows(matrix)
            if score > best_score or (not best_rows and rows):
                best_score, best_rows = score, rows
        return best_rows
    raise HTTPException(status_code=415, detail="Envie uma planilha CSV ou XLSX")


@router.post("/import", response_model=CatalogImportRead)
async def import_catalog(
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
    file: UploadFile = File(...),
    mapping_confirmed: bool = Form(False),
    name_column: str | None = Form(None),
    price_column: str | None = Form(None),
    category_column: str | None = Form(None),
    sku_column: str | None = Form(None),
) -> CatalogImportRead:
    filename = file.filename or ""
    if not filename.lower().endswith((".csv", ".xlsx")):
        raise HTTPException(status_code=415, detail="Envie uma planilha CSV ou XLSX")
    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="A planilha está vazia")
    if len(content) > MAX_CATALOG_FILE_BYTES:
        raise HTTPException(status_code=413, detail="A planilha deve ter no máximo 20 MB")
    try:
        raw_rows = spreadsheet_rows(filename, content)
    except UnicodeDecodeError:
        raise HTTPException(status_code=422, detail="Não foi possível ler o CSV. Salve-o como UTF-8 e tente novamente") from None
    if len(raw_rows) > MAX_CATALOG_ROWS:
        raise HTTPException(status_code=422, detail=f"A planilha pode ter no máximo {MAX_CATALOG_ROWS} itens")
    if not raw_rows:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Não encontramos uma tabela válida na planilha",
                "errors": [
                    "Use uma coluna de nome (Produto, Serviço, Item ou Descrição) e uma coluna de valor "
                    "(Preço, Valor, Preço de venda ou Valor unitário)."
                ],
            },
        )

    headers = list(raw_rows[0][1])
    suggested_mapping = {
        "name": matching_header(headers, "name"),
        "price": matching_header(headers, "price"),
        "category": matching_header(headers, "category"),
        "sku": matching_header(headers, "sku"),
    }
    if not mapping_confirmed and (not suggested_mapping["name"] or not suggested_mapping["price"]):
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Confirme quais colunas representam o item e o preço",
                "mapping_required": True,
                "headers": headers,
                "suggested_mapping": suggested_mapping,
            },
        )

    selected_mapping = {
        "name": normalize_header(name_column) if mapping_confirmed and name_column else None,
        "price": normalize_header(price_column) if mapping_confirmed and price_column else None,
        "category": normalize_header(category_column) if mapping_confirmed and category_column else None,
        "sku": normalize_header(sku_column) if mapping_confirmed and sku_column else None,
    }
    if mapping_confirmed and not selected_mapping["name"]:
        raise HTTPException(status_code=422, detail="Selecione a coluna que contém o nome do item")

    imported = updated = skipped = 0
    errors: list[str] = []
    price_headers = set(CATALOG_HEADER_ALIASES["price"])
    if not selected_mapping["price"] and not any(
        any(alias in normalize_header(header) for alias in price_headers) for header in headers
    ):
        errors.append(
            "Nenhuma coluna de preço foi reconhecida; os itens foram publicados como 'Sob consulta'."
        )
    results: list[CatalogItem] = []
    for row_number, raw_row in raw_rows:
        row = {normalize_header(key): value for key, value in raw_row.items()}
        name_value = row.get(selected_mapping["name"]) if selected_mapping["name"] else first_value(row, "name")
        name = str(name_value).strip() if name_value is not None else None
        if not name:
            skipped += 1
            errors.append(f"Linha {row_number}: informe o nome do produto ou serviço")
            continue
        price_value = row.get(selected_mapping["price"]) if selected_mapping["price"] else first_value(row, "price")
        price = parse_price(price_value)
        if price_value is not None and price is None:
            skipped += 1
            errors.append(f"Linha {row_number}: preço inválido")
            continue
        sku_value = row.get(selected_mapping["sku"]) if selected_mapping["sku"] else first_value(row, "sku")
        sku = str(sku_value).strip() if sku_value is not None else None
        item = None
        if sku:
            item = await db.scalar(select(CatalogItem).where(CatalogItem.tenant_id == tenant.id, CatalogItem.sku == sku))
        values = {
            "name": name[:200],
            "description": str(value).strip() if (value := first_value(row, "description")) is not None else None,
            "category": str(value).strip() if (value := (row.get(selected_mapping["category"]) if selected_mapping["category"] else first_value(row, "category"))) is not None else None,
            "sku": sku,
            "price": price,
            "currency": str(first_value(row, "currency") or "BRL").upper()[:3],
        }
        if item:
            for key, value in values.items():
                setattr(item, key, value)
            updated += 1
        else:
            item = CatalogItem(tenant_id=tenant.id, source="import", **values)
            db.add(item)
            imported += 1
        results.append(item)
    if not results:
        raise HTTPException(status_code=422, detail={"message": "Nenhum item válido foi encontrado", "errors": errors[:20]})
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Há itens duplicados na planilha") from exc
    for item in results:
        await db.refresh(item)
    return CatalogImportRead(imported=imported, updated=updated, skipped=skipped, errors=errors[:20], items=[CatalogItemRead.model_validate(item) for item in results])


@router.get("/items", response_model=list[CatalogItemRead])
async def list_catalog_items(tenant: CurrentTenant, db: DbSession) -> list[CatalogItemRead]:
    items = await db.scalars(
        select(CatalogItem)
        .where(CatalogItem.tenant_id == tenant.id)
        .order_by(CatalogItem.is_active.desc(), CatalogItem.name.asc())
    )
    return [CatalogItemRead.model_validate(item) for item in items]


@router.post("/items", response_model=CatalogItemRead, status_code=status.HTTP_201_CREATED)
async def create_catalog_item(
    payload: CatalogItemCreate,
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
) -> CatalogItemRead:
    values = payload.model_dump()
    values["currency"] = payload.currency.upper()
    item = CatalogItem(
        tenant_id=tenant.id,
        source="manual",
        **values,
    )
    db.add(item)
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Já existe um item com este SKU") from exc
    await db.refresh(item)
    return CatalogItemRead.model_validate(item)


@router.post("/items/publish")
async def publish_catalog_items(
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
) -> dict[str, int]:
    result = await db.execute(
        update(CatalogItem)
        .where(CatalogItem.tenant_id == tenant.id)
        .values(is_active=True)
    )
    await db.commit()
    return {"saved": result.rowcount or 0}


@router.delete("/imported-items")
async def delete_imported_catalog_items(
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
) -> dict[str, int]:
    result = await db.execute(
        delete(CatalogItem).where(
            CatalogItem.tenant_id == tenant.id,
            CatalogItem.source == "import",
        )
    )
    await db.commit()
    return {"deleted": result.rowcount or 0}


@router.delete("/items/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_catalog_item(
    item_id: uuid.UUID,
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
) -> Response:
    result = await db.execute(
        delete(CatalogItem).where(
            CatalogItem.id == item_id,
            CatalogItem.tenant_id == tenant.id,
        )
    )
    if not result.rowcount:
        await db.rollback()
        raise HTTPException(status_code=404, detail="Item não encontrado")
    await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.patch("/items/{item_id}", response_model=CatalogItemRead)
async def update_catalog_item(
    item_id: uuid.UUID,
    payload: CatalogItemUpdate,
    tenant: CurrentTenant,
    db: DbSession,
    _csrf: CsrfGuard,
) -> CatalogItemRead:
    item = await db.scalar(
        select(CatalogItem).where(CatalogItem.id == item_id, CatalogItem.tenant_id == tenant.id)
    )
    if not item:
        raise HTTPException(status_code=404, detail="Item não encontrado")
    values = payload.model_dump(exclude_unset=True)
    if values.get("currency"):
        values["currency"] = values["currency"].upper()
    for field, value in values.items():
        setattr(item, field, value)
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Já existe um item com este SKU") from exc
    await db.refresh(item)
    return CatalogItemRead.model_validate(item)
